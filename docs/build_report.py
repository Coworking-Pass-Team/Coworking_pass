"""Compiles the markdown documentation into one Word report per language.

Usage:  python docs/build_report.py          # builds both reports
        python docs/build_report.py en|ar    # builds a single one
Output: docs/Coworking_Pass_Final_Report_EN.docx    (from docs/*.md)
        docs/ar/Coworking_Pass_Final_Report_AR.docx (from docs/ar/*.md, kept local)
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from report_text import LANGS, FILES  # noqa: E402

LANG = LANGS['en']  # switched per build in main()

INK = RGBColor(0x2D, 0x35, 0x36)       # dark slate used across the product UI
ACCENT = RGBColor(0x40, 0x53, 0x4C)    # moss green
MUTED = RGBColor(0x6B, 0x77, 0x73)
LATIN_FONT = 'Calibri'
ARABIC_FONT = 'Arial'
MONO_FONT = 'Consolas'

ARABIC_RE = re.compile(r'[؀-ۿݐ-ݿ]')
EMOJI_RE = re.compile('[\U0001F000-\U0001FAFF☀-➿️‍]')


# ---------------------------------------------------------------- low-level helpers
def is_arabic(text: str) -> bool:
    letters = re.findall(r'[A-Za-z؀-ۿ]', text)
    if not letters:
        return False
    return sum(1 for c in letters if ARABIC_RE.match(c)) >= len(letters) / 2


def clean(text: str) -> str:
    """Removes emoji (they do not render reliably in Word) and collapses whitespace."""
    return re.sub(r'\s+', ' ', EMOJI_RE.sub('', text)).strip()


def set_run_fonts(run, mono=False, rtl=False):
    name = MONO_FONT if mono else (ARABIC_FONT if rtl else LATIN_FONT)
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(attr), name)
    if rtl:
        rpr.append(OxmlElement('w:rtl'))


def set_bidi(paragraph, rtl: bool):
    if not rtl:
        return
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(OxmlElement('w:bidi'))


def shade(element_pr, fill: str):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    element_pr.append(shd)


def paragraph_border(paragraph, side='left', color='9AA7A1', size=18):
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn('w:pBdr'))
    if borders is None:
        borders = OxmlElement('w:pBdr')
        ppr.append(borders)
    edge = OxmlElement('w:' + side)
    edge.set(qn('w:val'), 'single')
    edge.set(qn('w:sz'), str(size))
    edge.set(qn('w:space'), '8')
    edge.set(qn('w:color'), color)
    borders.append(edge)


INLINE_RE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))')


def add_inline(paragraph, text, rtl=False, size=None, color=None, bold=False, italic=False):
    """Adds text with **bold**, *italic*, `code` and [link](url) markup as separate runs."""
    for part in INLINE_RE.split(clean(text)):
        if not part:
            continue
        run_bold, run_italic, mono = bold, italic, False
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            part, run_bold = part[2:-2], True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            part, mono = part[1:-1], True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            part, run_italic = part[1:-1], True
        elif part.startswith('[') and '](' in part:
            part = part[1:part.index('](')]
        run = paragraph.add_run(part)
        run.bold = run_bold
        run.italic = run_italic
        set_run_fonts(run, mono=mono, rtl=rtl and not mono)
        if mono:
            run.font.size = Pt((size or 10.5) - 1)
            rpr = run._element.get_or_add_rPr()
            shade(rpr, 'EEF1EF')
        elif size:
            run.font.size = Pt(size)
        if color is not None and not mono:
            run.font.color.rgb = color


# ---------------------------------------------------------------- document setup
def configure_styles(doc: Document):
    normal = doc.styles['Normal']
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        fonts.set(qn(attr), LATIN_FONT)
    fonts.set(qn('w:cs'), ARABIC_FONT)

    sizes = {'Heading 1': 22, 'Heading 2': 16, 'Heading 3': 13, 'Heading 4': 11.5}
    for name, size in sizes.items():
        style = doc.styles[name]
        style.font.name = LATIN_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK if name == 'Heading 1' else ACCENT
        style.paragraph_format.space_before = Pt(18 if name == 'Heading 1' else 12)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
        srpr = style.element.get_or_add_rPr()
        sfonts = srpr.find(qn('w:rFonts'))
        if sfonts is None:
            sfonts = OxmlElement('w:rFonts')
            srpr.append(sfonts)
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            sfonts.set(qn(attr), LATIN_FONT if attr != 'w:cs' else ARABIC_FONT)

    for name in ('List Bullet', 'List Bullet 2', 'List Bullet 3'):
        doc.styles[name].font.name = LATIN_FONT
        doc.styles[name].paragraph_format.space_after = Pt(3)


def add_page_number_footer(section, label):
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    para.text = ''
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(label + '   |   ' + LANG['page'])
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    for kind, text in (('begin', None), (None, ' PAGE '), ('end', None)):
        r = para.add_run()
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        if kind:
            fld = OxmlElement('w:fldChar')
            fld.set(qn('w:fldCharType'), kind)
            r._r.append(fld)
        else:
            instr = OxmlElement('w:instrText')
            instr.set(qn('xml:space'), 'preserve')
            instr.text = text
            r._r.append(instr)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- cover + summary
def centered(doc, text, size, color=INK, bold=False, italic=False, before=0):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(before)
    rtl = is_arabic(text)
    set_bidi(para, rtl)
    run = para.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    set_run_fonts(run, rtl=rtl)
    return para


def build_cover(doc: Document):
    for _ in range(5):
        doc.add_paragraph()
    bar = doc.add_paragraph()
    bar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_border(bar, 'bottom', color='40534C', size=24)
    centered(doc, 'Coworking Pass', 44, INK, bold=True, before=24)
    centered(doc, LANG['subtitle'], 16, ACCENT, italic=True, before=6)
    if LANG['tagline']:
        centered(doc, LANG['tagline'], 14, MUTED, before=6)
    for _ in range(6):
        doc.add_paragraph()
    for label, value in LANG['labels']:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(2)
        rtl = is_arabic(label + value)
        set_bidi(para, rtl)
        a = para.add_run(label + ':  ')
        a.font.size = Pt(12)
        a.font.color.rgb = MUTED
        set_run_fonts(a, rtl=rtl)
        b = para.add_run(value)
        b.bold = True
        b.font.size = Pt(12)
        b.font.color.rgb = INK
        set_run_fonts(b, rtl=rtl)
    page_break(doc)


def heading(doc, text, level=1):
    para = doc.add_paragraph(style='Heading %d' % level)
    rtl = is_arabic(text)
    set_bidi(para, rtl)
    add_inline(para, text, rtl=rtl)
    return para


def body_paragraph(doc, text, style=None, size=None, color=None, italic=False):
    para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    rtl = is_arabic(text)
    set_bidi(para, rtl)
    add_inline(para, text, rtl=rtl, size=size, color=color, italic=italic)
    return para


def add_toc(doc: Document):
    heading(doc, LANG['toc_title'], 1)
    para = doc.add_paragraph()
    set_bidi(para, is_arabic(LANG['toc_hint']))
    run = para.add_run()
    for kind, text in (('begin', None), (None, ' TOC \\o "1-2" \\h \\z \\u '), ('separate', None), (None, None), ('end', None)):
        if kind:
            fld = OxmlElement('w:fldChar')
            fld.set(qn('w:fldCharType'), kind)
            run._r.append(fld)
        elif text:
            instr = OxmlElement('w:instrText')
            instr.set(qn('xml:space'), 'preserve')
            instr.text = text
            run._r.append(instr)
        else:
            t = OxmlElement('w:t')
            t.text = LANG['toc_hint']
            run._r.append(t)
    page_break(doc)


def build_executive_summary(doc: Document):
    heading(doc, LANG['exec_title'], 1)
    for text in LANG['exec_intro']:
        body_paragraph(doc, text)
    heading(doc, LANG['cap_title'], 2)
    for text in LANG['capabilities']:
        body_paragraph(doc, text, style='List Bullet')
    heading(doc, LANG['structure_title'], 2)
    for title in LANG['chapters']:
        body_paragraph(doc, title, style='List Bullet')
    body_paragraph(doc, LANG['structure_note'], size=10, color=MUTED, italic=True)
    page_break(doc)


# ---------------------------------------------------------------- markdown rendering
def add_code_block(doc, lines, language=''):
    label = LANG['code_labels'][0] if language == 'mermaid' else (LANG['code_labels'][1] if not language else language.upper())
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(0)
    cap.paragraph_format.keep_with_next = True
    add_inline(cap, label, size=8.5, color=MUTED, bold=True)
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.left_indent = Cm(0.3)
        ppr = para._p.get_or_add_pPr()
        shade(ppr, 'F2F4F3')
        run = para.add_run(line.rstrip() if line.strip() else ' ')
        set_run_fonts(run, mono=True)
        run.font.size = Pt(8.5)
        run.font.color.rgb = INK
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def set_cell_background(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def add_table(doc, rows):
    header = rows[0]
    body = rows[1:]
    cols = len(header)
    table = doc.add_table(rows=1 + len(body), cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    arabic_table = sum(is_arabic(c) for r in rows for c in r) > (len(rows) * cols) / 2
    if arabic_table:
        tblpr = table._tbl.tblPr
        tblpr.append(OxmlElement('w:bidiVisual'))
    for r_index, row in enumerate([header] + body):
        for c_index in range(cols):
            text = row[c_index] if c_index < len(row) else ''
            cell = table.rows[r_index].cells[c_index]
            cell.text = ''
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            rtl = is_arabic(text)
            set_bidi(para, rtl)
            if r_index == 0:
                set_cell_background(cell, '40534C')
                add_inline(para, text, rtl=rtl, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
            else:
                if r_index % 2 == 0:
                    set_cell_background(cell, 'F4F6F5')
                add_inline(para, text, rtl=rtl, size=10)
    # repeat header row on each page
    trpr = table.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def split_row(line):
    cells = line.strip().strip('|').split('|')
    return [c.strip() for c in cells]


def render_markdown(doc: Document, text: str, chapter_title: str):
    lines = text.splitlines()
    i = 0
    first_h1_skipped = False
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped or re.fullmatch(r'-{3,}', stripped):
            i += 1
            continue

        if stripped.startswith('```'):
            language = stripped[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                block.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, block, language)
            continue

        heading = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if heading:
            level = len(heading.group(1))
            title = clean(heading.group(2))
            if level == 1 and not first_h1_skipped:
                first_h1_skipped = True  # the chapter heading replaces the file title
                sub = doc.add_paragraph()
                sub.paragraph_format.space_after = Pt(10)
                set_bidi(sub, is_arabic(title))
                add_inline(sub, title, rtl=is_arabic(title), size=12, color=MUTED, italic=True)
            else:
                mapped = min(max(level, 2) if level > 1 else 2, 4)
                para = doc.add_paragraph(style='Heading %d' % (mapped if level > 1 else 2))
                rtl = is_arabic(title)
                set_bidi(para, rtl)
                add_inline(para, title, rtl=rtl)
                for run in para.runs:
                    run.font.color.rgb = INK if mapped == 2 else ACCENT
                    run.bold = True
            i += 1
            continue

        if stripped.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = split_row(lines[i])
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in row if c):
                    rows.append(row)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        if stripped.startswith('>'):
            quote = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                quote.append(lines[i].strip().lstrip('>').strip())
                i += 1
            para = doc.add_paragraph()
            body = ' '.join(q for q in quote if q)
            rtl = is_arabic(body)
            set_bidi(para, rtl)
            para.paragraph_format.left_indent = Cm(0.5)
            paragraph_border(para, 'right' if rtl else 'left', color='40534C', size=24)
            shade(para._p.get_or_add_pPr(), 'EEF3F0')
            add_inline(para, body, rtl=rtl, size=10.5)
            continue

        bullet = re.match(r'^(\s*)([-*+])\s+(.*)$', line)
        numbered = re.match(r'^(\s*)(\d+)[.)]\s+(.*)$', line)
        if bullet or numbered:
            m = bullet or numbered
            indent = len(m.group(1).replace('\t', '    '))
            level = min(indent // 2, 2)
            body = m.group(3)
            # continuation lines belonging to the same item (indented text without a marker)
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r'^\s*([-*+]|\d+[.)])\s+', lines[i + 1]) \
                    and not lines[i + 1].strip().startswith(('#', '|', '>', '```')) and lines[i + 1].startswith('  '):
                i += 1
                body += ' ' + lines[i].strip()
            rtl = is_arabic(body)
            if bullet:
                para = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet %d' % (level + 1))
            else:
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(0.9 + level * 0.7)
                para.paragraph_format.first_line_indent = Cm(-0.6)
                para.paragraph_format.space_after = Pt(3)
                add_inline(para, m.group(2) + '.  ', size=None, bold=True)
            set_bidi(para, rtl)
            add_inline(para, body, rtl=rtl)
            i += 1
            continue

        # plain paragraph: merge following soft-wrapped lines
        body = stripped
        while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r'^(#{1,6}\s|\||>|```|\s*([-*+]|\d+[.)])\s|-{3,}$)', lines[i + 1].strip()):
            i += 1
            body += ' ' + lines[i].strip()
        rtl = is_arabic(body)
        para = doc.add_paragraph()
        set_bidi(para, rtl)
        add_inline(para, body, rtl=rtl)
        i += 1


# ---------------------------------------------------------------- schema order fix-up
PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs',
             'suppressAutoHyphens', 'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
             'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc', 'textDirection', 'textAlignment',
             'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange']
RPR_ORDER = ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike', 'outline', 'shadow', 'emboss', 'imprint', 'noProof',
             'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing', 'w', 'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd',
             'fitText', 'vertAlign', 'rtl', 'cs', 'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath']
TBLPR_ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd',
               'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook']


def reorder(element, order):
    def key(child):
        name = child.tag.split('}')[-1]
        return order.index(name) if name in order else len(order)
    children = sorted(list(element), key=key)
    for child in list(element):
        element.remove(child)
    for child in children:
        element.append(child)


def fix_schema_order(doc):
    parts = [doc.element.body] + [s.footer._element for s in doc.sections]
    for root in parts:
        for tag, order in (('w:pPr', PPR_ORDER), ('w:rPr', RPR_ORDER), ('w:tblPr', TBLPR_ORDER)):
            for el in root.iter(qn(tag)):
                reorder(el, order)


# ---------------------------------------------------------------- main
def build(lang_key: str):
    global LANG
    LANG = LANGS[lang_key]
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.0)
    configure_styles(doc)

    core = doc.core_properties
    core.title = 'Coworking Pass - Comprehensive Engineering, Architecture & Requirements Specification Report'
    core.subject = 'Final project report, version 2.0'
    core.author = 'Coworking Pass Team'
    core.language = 'ar-SA' if lang_key == 'ar' else 'en-US'

    build_cover(doc)
    add_page_number_footer(section, LANG['footer'])
    add_toc(doc)
    build_executive_summary(doc)

    for index, (title, filename) in enumerate(zip(LANG['chapters'], FILES)):
        path = LANG['source'] / filename
        text = path.read_text(encoding='utf-8')
        heading(doc, title, 1)
        render_markdown(doc, text, title)
        if index < len(FILES) - 1:
            page_break(doc)
        print('  [%s] added %s (%d lines)' % (lang_key, path.name, len(text.splitlines())))

    fix_schema_order(doc)
    doc.save(LANG['output'])
    print('saved', LANG['output'])


def main():
    targets = [a for a in sys.argv[1:] if a in LANGS] or list(LANGS)
    for key in targets:
        build(key)


if __name__ == '__main__':
    main()
